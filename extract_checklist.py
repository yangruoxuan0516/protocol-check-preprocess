#!/usr/bin/env python3
"""Extract Correctness and Completeness checklist rows from an XLSX file.

Output is UTF-8 JSON Lines (JSONL): one JSON object per title or checklist
item.  Only the sheets named ``Correctness Checklist`` and
``Completeness Checklist`` are read.  Rows 1-3 and columns Pass/Fail,
Comments, and Checker are intentionally ignored.

Example
-------
    python extract_checklist.py checklist.xlsx
    python extract_checklist.py checklist.xlsx -o checklist.jsonl

Dependency: openpyxl (``pip install openpyxl``)
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


HEADER_ROW = 4
FIRST_DATA_ROW = 5
EXPECTED_HEADERS = (
    "Number",
    "List Value",
    "Description",
    "Reference",
    "Applicability",
    "Pass/Fail",
    "Comments",
    "Checker",
)
SHEET_SPECS = (
    ("Correctness Checklist", "correctness", "COR"),
    ("Completeness Checklist", "completeness", "CPL"),
)


class ChecklistExtractionError(RuntimeError):
    """Raised when the workbook does not match the expected checklist format."""


@dataclass(frozen=True)
class ParsedReference:
    document: str
    target: str
    canonical: str


@dataclass
class ChecklistRecord:
    source_file: str
    sheet: str
    checklist_type: str
    excel_row: int
    number: str
    row_type: str
    parent_number: str | None
    parent_title: str | None
    list_value: str | None
    description: str | None
    reference_raw: str | None
    references: list[dict[str, str]]
    reference_unparsed: list[str]
    applicability: str | None


REFERENCE_RE = re.compile(
    r"(?ix)"
    r"(?P<arp>\bARP\s*4754[A-Z]?\s+"
    r"(?P<arp_section>\d+(?:\.\d+)+)"
    r"(?:\s+(?P<arp_item>[a-z])"
    r"(?P<arp_subitems>(?:\s*\(\d+\)(?:\s*,?\s*\(\d+\))*)?)"
    r"(?![a-z0-9(]))?)"
    r"|"
    r"(?P<std>\bENG[\s-]*STD[\s-]*011\s+RS[\s-]*"
    r"(?P<std_number>\d+))"
    r"|"
    r"(?P<do>\bDO[\s-]*297\s+"
    r"(?P<do_section>\d+(?:\.\d+)+)"
    r"(?:\s+(?P<do_item>[a-z])(?![a-z0-9(]))?)"
)

REFERENCE_FILLER_RE = re.compile(
    r"(?ix)^"
    r"(?:"
    r"\s+|[,;:/&+\-\u2013\u2014]+|"
    r"\b(?:and/or|and|or|see|also|refer(?:ring)?\s+to)\b"
    r")*"
    r"$"
)


def _clean_cell(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return text or None


def _normalize_header(value: object) -> str:
    text = _clean_cell(value) or ""
    return re.sub(r"\s+", " ", text).casefold()


def parse_reference_cell(
    reference_raw: str | None,
) -> tuple[list[dict[str, str]], list[str]]:
    """Return parsed references and nontrivial unmatched source fragments."""
    if not reference_raw:
        return [], []

    parsed: list[ParsedReference] = []
    seen: set[str] = set()
    matches = list(REFERENCE_RE.finditer(reference_raw))
    for match in matches:
        if match.group("arp"):
            section = match.group("arp_section")
            item = match.group("arp_item")
            subitems = re.findall(r"\((\d+)\)", match.group("arp_subitems") or "")
            if item and subitems:
                targets = [
                    f"{section} {item.lower()}({subitem})"
                    for subitem in subitems
                ]
            else:
                targets = [section + (f" {item.lower()}" if item else "")]
            results = [
                ParsedReference(
                    document="ARP4754",
                    target=target,
                    canonical=f"ARP4754 {target}",
                )
                for target in targets
            ]
        elif match.group("std"):
            target = f"RS{int(match.group('std_number'))}"
            results = [
                ParsedReference(
                    document="ENG-STD-011",
                    target=target,
                    canonical=f"ENG-STD-011 {target}",
                )
            ]
        else:
            section = match.group("do_section")
            item = match.group("do_item")
            target = section + (f" {item.lower()}" if item else "")
            results = [
                ParsedReference(
                    document="DO-297",
                    target=target,
                    canonical=f"DO-297 {target}",
                )
            ]

        for result in results:
            if result.canonical not in seen:
                parsed.append(result)
                seen.add(result.canonical)

    unparsed: list[str] = []
    cursor = 0
    for match in matches:
        fragment = reference_raw[cursor:match.start()]
        if fragment and not REFERENCE_FILLER_RE.fullmatch(fragment):
            cleaned = fragment.strip(" \t\r\n,;:/&+-\u2013\u2014")
            if cleaned:
                unparsed.append(cleaned)
        cursor = match.end()
    fragment = reference_raw[cursor:]
    if fragment and not REFERENCE_FILLER_RE.fullmatch(fragment):
        cleaned = fragment.strip(" \t\r\n,;:/&+-\u2013\u2014")
        if cleaned:
            unparsed.append(cleaned)

    return [asdict(reference) for reference in parsed], unparsed


def parse_references(reference_raw: str | None) -> list[dict[str, str]]:
    """Find and canonicalize supported references while preserving their order."""
    references, _unparsed = parse_reference_cell(reference_raw)
    return references


def _resolve_sheet(workbook, expected_name: str) -> Worksheet:
    matches = {
        name.strip().casefold(): name
        for name in workbook.sheetnames
    }
    actual_name = matches.get(expected_name.casefold())
    if actual_name is None:
        raise ChecklistExtractionError(
            f"Required sheet {expected_name!r} was not found. "
            f"Available sheets: {workbook.sheetnames}"
        )
    return workbook[actual_name]


def _validate_headers(sheet: Worksheet) -> None:
    actual = tuple(
        _normalize_header(sheet.cell(HEADER_ROW, column).value)
        for column in range(1, 9)
    )
    expected = tuple(_normalize_header(header) for header in EXPECTED_HEADERS)
    if actual != expected:
        shown = [sheet.cell(HEADER_ROW, column).value for column in range(1, 9)]
        raise ChecklistExtractionError(
            f"Unexpected headers in {sheet.title!r} row {HEADER_ROW}: {shown}. "
            f"Expected: {list(EXPECTED_HEADERS)}"
        )


def _number_pattern(prefix: str) -> re.Pattern[str]:
    return re.compile(rf"^{re.escape(prefix)}-\d{{2}}(?:\.\d{{2}})?$")


def _extract_sheet_records(
    sheet: Worksheet,
    *,
    source_file: str,
    checklist_type: str,
    prefix: str,
    allow_invalid_rows: bool,
) -> list[ChecklistRecord]:
    _validate_headers(sheet)
    number_re = _number_pattern(prefix)
    raw_records: list[ChecklistRecord] = []
    seen_numbers: set[str] = set()

    rows = sheet.iter_rows(
        min_row=FIRST_DATA_ROW,
        max_col=8,
        values_only=True,
    )
    for row_number, row in enumerate(rows, start=FIRST_DATA_ROW):
        values = [_clean_cell(value) for value in row]
        if not any(values):
            continue

        number = (values[0] or "").upper()
        if not number_re.fullmatch(number):
            message = (
                f"{sheet.title}!A{row_number} has invalid Number "
                f"{values[0]!r}; expected {prefix}-NN or {prefix}-NN.NN."
            )
            if allow_invalid_rows:
                print(f"Warning: {message} Row skipped.", file=sys.stderr)
                continue
            raise ChecklistExtractionError(message)
        if number in seen_numbers:
            raise ChecklistExtractionError(
                f"Duplicate Number {number!r} in {sheet.title!r}."
            )
        seen_numbers.add(number)

        is_title = "." not in number
        parent_number = None if is_title else number.split(".", 1)[0]
        references, reference_unparsed = parse_reference_cell(values[3])
        raw_records.append(
            ChecklistRecord(
                source_file=source_file,
                sheet=sheet.title,
                checklist_type=checklist_type,
                excel_row=row_number,
                number=number,
                row_type="title" if is_title else "item",
                parent_number=parent_number,
                parent_title=None,
                list_value=values[1],
                description=values[2],
                reference_raw=values[3],
                references=references,
                reference_unparsed=reference_unparsed,
                applicability=values[4],
            )
        )

    title_map = {
        record.number: record.list_value
        for record in raw_records
        if record.row_type == "title"
    }
    for record in raw_records:
        if record.parent_number is not None:
            record.parent_title = title_map.get(record.parent_number)
            if record.parent_title is None:
                message = (
                    f"{record.number!r} in {sheet.title!r} has no matching "
                    f"title row {record.parent_number!r}."
                )
                if allow_invalid_rows:
                    print(f"Warning: {message}", file=sys.stderr)
                else:
                    raise ChecklistExtractionError(message)
    return raw_records


def extract_checklist(
    workbook_path: str | Path,
    *,
    allow_invalid_rows: bool = False,
) -> list[ChecklistRecord]:
    """Read the two checklist sheets and return records in workbook order."""
    path = Path(workbook_path)
    if path.suffix.casefold() != ".xlsx":
        raise ValueError(f"Expected an .xlsx file, got {path.suffix or 'no extension'!r}.")
    if not path.is_file():
        raise FileNotFoundError(path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        records: list[ChecklistRecord] = []
        for sheet_name, checklist_type, prefix in SHEET_SPECS:
            sheet = _resolve_sheet(workbook, sheet_name)
            records.extend(
                _extract_sheet_records(
                    sheet,
                    source_file=path.name,
                    checklist_type=checklist_type,
                    prefix=prefix,
                    allow_invalid_rows=allow_invalid_rows,
                )
            )
        return records
    finally:
        workbook.close()


def write_jsonl(records: Sequence[ChecklistRecord], output_path: str | Path) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for record in records:
            stream.write(json.dumps(asdict(record), ensure_ascii=False) + "\n")
    return path


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Extract Correctness and Completeness checklist rows from XLSX "
            "into flat JSONL."
        )
    )
    parser.add_argument("workbook", type=Path, help="path to checklist.xlsx")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="output JSONL path (default: <workbook>_checklist.jsonl)",
    )
    parser.add_argument(
        "--allow-invalid-rows",
        action="store_true",
        help="skip invalid numbered rows and warn instead of stopping",
    )
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    output = args.output or args.workbook.with_name(
        f"{args.workbook.stem}_checklist.jsonl"
    )
    try:
        records = extract_checklist(
            args.workbook,
            allow_invalid_rows=args.allow_invalid_rows,
        )
        output = write_jsonl(records, output)
    except (
        ChecklistExtractionError,
        FileNotFoundError,
        OSError,
        ValueError,
    ) as exc:
        raise SystemExit(f"Error: {exc}") from exc

    titles = sum(record.row_type == "title" for record in records)
    items = len(records) - titles
    print(
        f"Wrote {len(records)} records ({titles} titles, {items} items) "
        f"to {output}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())