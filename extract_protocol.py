#!/usr/bin/env python3
"""Extract a four-column NIO table from a DOCX file.

Expected columns, by position:
    1. ID (NIO-<digits>)
    2. Specification
    3. Rationale (may be empty)
    4. Req (Yes/No, or empty)

Outputs:
    <document_stem>_records.jsonl  - convenient for Python/RAG pipelines
    <document_stem>_records.csv    - convenient for manual review in Excel
    <document_stem>_summary.json   - extraction counts and validation warnings

Usage:
    python docx_table_to_records.py protocol.docx
    python docx_table_to_records.py protocol.docx --output-dir extracted
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path
from typing import Any

from docx import Document


ID_PATTERN = re.compile(r"^NIO-\d+$", re.IGNORECASE)
HEADER_ID_NAMES = {"id", "identifier", "nio id", "nio-id"}


def clean_cell_text(text: str) -> str:
    """Normalize Word whitespace while retaining meaningful line breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    lines = []
    for line in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def canonical_req(value: str) -> str | None:
    value_lower = value.strip().lower()
    if not value_lower:
        return None
    if value_lower == "yes":
        return "Yes"
    if value_lower == "no":
        return "No"
    return value.strip()


def is_header_row(first_cell: str) -> bool:
    return first_cell.strip().lower() in HEADER_ID_NAMES


def extract_records(docx_path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    document = Document(docx_path)
    records: list[dict[str, Any]] = []
    warnings: list[str] = []
    seen_ids: set[str] = set()

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [clean_cell_text(cell.text) for cell in row.cells]

            if not any(cells):
                continue

            if is_header_row(cells[0]):
                continue

            if len(cells) != 4:
                warnings.append(
                    f"Table {table_index}, row {row_index}: expected 4 columns, "
                    f"found {len(cells)}; row skipped."
                )
                continue

            nio_id, specification, rationale, req = cells
            nio_id = nio_id.upper()
            req = canonical_req(req)
            row_errors: list[str] = []

            if not ID_PATTERN.fullmatch(nio_id):
                row_errors.append(f"unexpected ID {nio_id!r}")
            if not specification:
                row_errors.append("empty specification")
            # An empty Req is valid and is stored as null. Only a non-empty,
            # unexpected value is reported as a warning.
            if req is not None and req not in {"Yes", "No"}:
                row_errors.append(f"Req is not Yes/No: {req!r}")
            if nio_id in seen_ids:
                row_errors.append(f"duplicate ID {nio_id!r}")

            if nio_id:
                seen_ids.add(nio_id)

            if row_errors:
                warnings.append(
                    f"Table {table_index}, row {row_index}: " + "; ".join(row_errors)
                )

            records.append(
                {
                    "id": nio_id,
                    "specification": specification,
                    "rationale": rationale or None,
                    "req": req,
                    "source": {
                        "file": docx_path.name,
                        "table": table_index,
                        "word_row": row_index,
                    },
                }
            )

    if not document.tables:
        warnings.append("The document contains no tables.")

    return records, warnings


def write_jsonl(records: list[dict[str, Any]], output_path: Path) -> None:
    with output_path.open("w", encoding="utf-8", newline="\n") as output_file:
        for record in records:
            output_file.write(json.dumps(record, ensure_ascii=False) + "\n")


def write_csv(records: list[dict[str, Any]], output_path: Path) -> None:
    fieldnames = [
        "id",
        "specification",
        "rationale",
        "req",
        "source_file",
        "source_table",
        "source_word_row",
    ]
    # utf-8-sig lets Windows Excel recognize UTF-8 reliably.
    with output_path.open("w", encoding="utf-8-sig", newline="") as output_file:
        writer = csv.DictWriter(output_file, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            writer.writerow(
                {
                    "id": record["id"],
                    "specification": record["specification"],
                    "rationale": record["rationale"] or "",
                    "req": record["req"] or "",
                    "source_file": record["source"]["file"],
                    "source_table": record["source"]["table"],
                    "source_word_row": record["source"]["word_row"],
                }
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extract a four-column NIO table from a DOCX file."
    )
    parser.add_argument("docx", type=Path, help="Path to the input .docx file")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Output directory (default: next to the input document)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    docx_path: Path = args.docx.expanduser().resolve()

    if not docx_path.is_file():
        print(f"Error: file not found: {docx_path}", file=sys.stderr)
        return 2
    if docx_path.suffix.lower() != ".docx":
        print("Error: input file must have a .docx extension", file=sys.stderr)
        return 2

    output_dir = (
        args.output_dir.expanduser().resolve()
        if args.output_dir is not None
        else docx_path.parent
    )
    output_dir.mkdir(parents=True, exist_ok=True)

    records, warnings = extract_records(docx_path)
    base_name = docx_path.stem
    jsonl_path = output_dir / f"{base_name}_records.jsonl"
    csv_path = output_dir / f"{base_name}_records.csv"
    summary_path = output_dir / f"{base_name}_summary.json"

    write_jsonl(records, jsonl_path)
    write_csv(records, csv_path)

    req_counts = {
        "Yes": sum(record["req"] == "Yes" for record in records),
        "No": sum(record["req"] == "No" for record in records),
        "empty": sum(record["req"] is None for record in records),
        "other": sum(
            record["req"] is not None and record["req"] not in {"Yes", "No"}
            for record in records
        ),
    }

    summary = {
        "input_file": str(docx_path),
        "record_count": len(records),
        "req_counts": req_counts,
        "warning_count": len(warnings),
        "warnings": warnings,
        "outputs": {
            "jsonl": str(jsonl_path),
            "csv": str(csv_path),
        },
    }
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"Extracted {len(records)} records.")
    print(f"JSONL:  {jsonl_path}")
    print(f"CSV:    {csv_path}")
    print(f"Summary:{summary_path}")
    if warnings:
        print(f"Warnings: {len(warnings)} (see the summary file)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())